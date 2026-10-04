"""Generate a labelled SYNTHETIC resume set (TXT / DOCX / single-column PDF / two-column PDF).

Gold labels come for free because every resume is rendered from a structured record.
Gold months are computed here independently of the project's own date code (no circularity).

  python -m tools.generate_dataset --n 40 --out data/synthetic --seed 42

Realism knobs: varied headings/section order, 4 experience layouts, 3 education layouts, 4 date formats,
skill aliases (sklearn, JS, k8s), skills OUTSIDE the taxonomy (Figma, Terraform ...), concurrent roles,
wrapped lines in two-column PDFs. Be transparent in your report that this set is synthetic.
"""
from __future__ import annotations

import argparse
import json
import random
from pathlib import Path

FIRST = ["Aarav", "Diya", "Kabir", "Meera", "Rohan", "Sana", "Vihaan", "Ishita", "Arnav", "Tanvi", "Nikhil", "Pooja",
         "Rahul", "Anjali", "Yash", "Kavya", "Dev", "Nisha", "Karan", "Simran"]
LAST = ["Kapoor", "Nair", "Chauhan", "Iyer", "Bhatt", "Joshi", "Reddy", "Saxena", "Menon", "Rathore", "Pillai",
        "Tiwari", "Desai", "Khanna", "Banerjee", "Shekhawat", "Gill", "Rao"]
CITIES = [("Jaipur", "Rajasthan"), ("Pune", "Maharashtra"), ("Bengaluru", "Karnataka"), ("Hyderabad", "Telangana"),
          ("Indore", "Madhya Pradesh"), ("Lucknow", "Uttar Pradesh"), ("Chennai", "Tamil Nadu"), ("Kochi", "Kerala")]
COLLEGES = ["Aravali Institute of Technology", "Deccan University", "Narmada College of Engineering",
            "Sahyadri Institute of Science", "Kaveri University", "Himalaya School of Computing",
            "Vindhya Institute of Technology", "Konkan University"]
SCHOOLS = ["Sunrise Public School", "Greenfield Senior School", "Lotus Valley School", "Orchid International School"]
DEGREES = ["B.Tech in Computer Science", "B.Tech in Information Technology", "B.E. in Electronics",
           "BCA", "B.Sc. in Data Science", "M.Tech in Software Engineering", "MCA", "M.Sc. in Mathematics"]
COMPANIES = ["Nimbus Analytics", "Orbit Labs", "Zenith Soft", "Kiran Labs", "Finlytics", "Quanta Works",
             "BlueLeaf Systems", "Helix Digital", "Pragati Tech", "Mosaic Data"]
ROLES = ["Data Science Intern", "Software Engineer Intern", "Backend Developer Intern", "Data Analyst Intern",
         "Research Intern", "QA Engineer Intern", "Web Developer", "Junior Analyst", "Project Trainee"]
PROJECTS = [("Library Management App", "Manages book issue and return records."),
            ("Sentiment Analyzer", "Classifies product reviews as positive or negative."),
            ("Expense Tracker", "Tracks daily spending with monthly charts."),
            ("Chat Support Bot", "Answers common student queries."),
            ("Weather Dashboard", "Shows live forecasts for selected cities."),
            ("Attendance Portal", "Marks and reports class attendance."),
            ("Price Predictor", "Estimates used-car prices from listings.")]
# canonical taxonomy skills -> forms they may be written as
KNOWN = {"Python": ["Python"], "SQL": ["SQL"], "Java": ["Java"], "C++": ["C++"], "JavaScript": ["JavaScript", "JS"],
         "React": ["React"], "Node.js": ["Node.js"], "Flask": ["Flask"], "Django": ["Django"], "FastAPI": ["FastAPI"],
         "Pandas": ["Pandas"], "NumPy": ["NumPy"], "scikit-learn": ["scikit-learn", "sklearn"],
         "TensorFlow": ["TensorFlow"], "PyTorch": ["PyTorch"], "spaCy": ["spaCy"], "Docker": ["Docker"],
         "Git": ["Git"], "Linux": ["Linux"], "AWS": ["AWS"], "Tableau": ["Tableau"], "Power BI": ["Power BI"],
         "Excel": ["Excel"], "MongoDB": ["MongoDB"], "PostgreSQL": ["PostgreSQL", "Postgres"],
         "Kubernetes": ["Kubernetes", "k8s"], "Hadoop": ["Hadoop"]}
UNKNOWN = ["Figma", "Selenium", "Terraform", "Elasticsearch", "Power Apps", "Looker", "Snowflake", "dbt",
           "Prometheus", "GraphQL"]
BULLETS = ["Built dashboards using {a} and {b}.", "Automated reporting pipelines with {a}.",
           "Developed internal services in {a}.", "Improved query speed in {a} by 30%.",
           "Deployed prototypes using {a} and {b}.", "Wrote unit tests and documentation.",
           "Worked with a team of 5 engineers.", "Cleaned and analysed datasets with {a}."]
HEAD = {"summary": ["SUMMARY", "Profile", "Objective"],
        "education": ["EDUCATION", "Education", "Academic Background", "Educational Qualifications"],
        "experience": ["EXPERIENCE", "Work Experience", "Professional Experience", "Internships"],
        "skills": ["SKILLS", "Technical Skills", "Technical Proficiency", "Key Skills"],
        "projects": ["PROJECTS", "Academic Projects", "Personal Projects"]}
MON = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
FULL = ["January", "February", "March", "April", "May", "June", "July", "August", "September", "October",
        "November", "December"]
LAYOUTS = ["txt", "docx", "pdf_single", "pdf_two_col"]


def fmt(idx: int, style: int) -> str:
    y, m = idx // 12, idx % 12
    return [f"{MON[m]} {y}", f"{FULL[m]} {y}", f"{m + 1:02d}/{y}", f"{MON[m]} '{str(y)[2:]}"][style]


def make_record(rng: random.Random, layout: str):
    first, last = rng.choice(FIRST), rng.choice(LAST)
    city, state = rng.choice(CITIES)
    slug = f"{first}-{last}".lower()
    email = f"{first}.{last}{rng.randint(1, 99)}@example.com".lower()
    pnum = f"9{rng.randint(10 ** 8, 10 ** 9 - 1)}"
    phone = rng.choice([f"+91 {pnum[:5]} {pnum[5:]}", f"+91-{pnum}", f"0{pnum}", pnum])

    # ---- skills
    known = rng.sample(list(KNOWN), rng.randint(6, 11))
    unknown = rng.sample(UNKNOWN, rng.randint(0, 3))
    written = {c: rng.choice(KNOWN[c]) for c in known} | {u: u for u in unknown}
    allsk = list(written)
    rng.shuffle(allsk)
    gold_skills = list(allsk)

    # ---- education
    dstyle = rng.randint(0, 2)
    deg, college = rng.choice(DEGREES), rng.choice(COLLEGES)
    y2 = rng.randint(2024, 2026)
    y1 = y2 - (2 if deg.startswith("M") else 3 if "BCA" in deg or "Sc" in deg else 4)
    score = f"{rng.randint(65, 94) / 10:.1f}"
    edu_lines, gold_edu = [], [{"degree": deg, "institute": college, "year": str(y2)}]
    if dstyle == 0:
        edu_lines += [("t", deg), ("t", f"{college} | {y1} - {y2}"), ("t", f"CGPA: {score}/10")]
    elif dstyle == 1:
        edu_lines += [("t", f"{deg}, {college} ({y1} - {y2}), CGPA {score}")]
    else:
        edu_lines += [("t", f"{deg} - {college} - {y2}")]
    if rng.random() < 0.5:
        school, yr = rng.choice(SCHOOLS), y1 - 1
        edu_lines.append(("t", f"Senior Secondary (Class XII), {school}, {city} | {yr}"))
        edu_lines.append(("t", f"Percentage: {rng.randint(70, 96)}%"))
        gold_edu.append({"degree": "Senior Secondary (Class XII)", "institute": school, "year": str(yr)})

    # ---- experience (month resolution; sometimes one concurrent role)
    n_exp = rng.choice([1, 2, 2, 3])
    estyle, datestyle, sep = rng.randint(0, 3), rng.randint(0, 3), rng.choice([" - ", " – ", " to "])
    cursor = rng.randint(2022, 2024) * 12 + rng.randint(0, 11)
    spans, exp_lines, gold_exp, covered = [], [], [], set()
    companies = rng.sample(COMPANIES, n_exp)
    for i, comp in enumerate(companies):
        dur = rng.randint(2, 8)
        start = cursor - rng.randint(1, dur - 1) if (i > 0 and rng.random() < 0.15) else cursor
        end = start + dur - 1
        cursor = max(cursor, end + 1) + rng.randint(0, 3)
        spans.append((start, end))
        covered |= set(range(start, end + 1))
        role = rng.choice(ROLES)
        rng_s = f"{fmt(start, datestyle)}{sep}{fmt(end, datestyle)}"
        ccity = rng.choice(CITIES)[0]
        if estyle == 0:
            exp_lines.append(("t", f"{role} | {comp}, {ccity} | {rng_s}"))
        elif estyle == 1:
            exp_lines.append(("t", f"{role} at {comp} ({rng_s})"))
        elif estyle == 2:
            exp_lines += [("t", f"{comp}, {ccity}"), ("t", f"{role}  {rng_s}")]
        else:
            exp_lines.append(("t", f"{comp} - {role} - {rng_s}"))
        a, b = rng.sample(allsk, 2)
        for tpl in rng.sample(BULLETS, rng.randint(1, 3)):
            exp_lines.append(("t", "- " + tpl.format(a=written[a], b=written[b])))
        gold_exp.append({"company": comp, "role": role})

    # ---- skills section
    sk_forms = [written[s] for s in allsk]
    sstyle = rng.randint(0, 3)
    if sstyle == 0:
        k = max(2, len(sk_forms) // 3)
        chunks = [sk_forms[i:i + k] for i in range(0, len(sk_forms), k)]
        sk_lines = [("t", f"{lab}: {', '.join(ch)}") for lab, ch in zip(["Languages", "Frameworks", "Tools", "Other"], chunks)]
    elif sstyle == 1:
        sk_lines = [("t", " | ".join(sk_forms))]
    elif sstyle == 2:
        sk_lines = [("t", ", ".join(sk_forms))]
    else:
        sk_lines = [("t", "- " + s) for s in sk_forms]

    # ---- projects
    proj_lines = []
    for name, desc in rng.sample(PROJECTS, rng.randint(1, 2)):
        tech = ", ".join(written[s] for s in rng.sample(allsk, 2))
        proj_lines += [("t", f"{name} | {tech}"), ("t", f"- {desc}")]

    # ---- header
    shown_name = f"{first} {last}".upper() if rng.random() < 0.3 else f"{first} {last}"
    multi = layout == "pdf_two_col" or rng.random() < 0.4
    header = [("name", shown_name)]
    if multi:
        header += [("t", f"{city}, {state}"), ("t", email), ("t", phone)]
    else:
        header.append(("t", f"{city}, {state} | {email} | {phone}"))
    if rng.random() < 0.7:
        header.append(("t", f"linkedin.com/in/{slug}"))
        header.append(("t", f"github.com/{slug.replace('-', '')}"))

    blocks = {"header": header,
              "education": [("h", rng.choice(HEAD["education"]))] + edu_lines,
              "experience": [("h", rng.choice(HEAD["experience"]))] + exp_lines,
              "skills": [("h", rng.choice(HEAD["skills"]))] + sk_lines,
              "projects": [("h", rng.choice(HEAD["projects"]))] + proj_lines}
    if rng.random() < 0.5:
        blocks["summary"] = [("h", rng.choice(HEAD["summary"])),
                             ("t", "Motivated student looking for roles in software and data.")]
    order = rng.choice([["education", "experience", "skills", "projects"], ["skills", "experience", "projects", "education"]])
    if "summary" in blocks:
        order.insert(0, "summary")
    gold = {"name": f"{first} {last}", "email": email, "phone": pnum, "skills": gold_skills, "education": gold_edu,
            "experience": gold_exp, "total_experience_months": len(covered), "layout": layout, "source": "synthetic"}
    return gold, blocks, order


# ------------------------------------------------------------------ renderers
def as_text(blocks, order) -> str:
    lines = []
    for sec in ["header"] + order:
        for kind, t in blocks[sec]:
            lines.append(t)
        lines.append("")
    return "\n".join(lines)


def write_docx(path, blocks, order):
    from docx import Document

    doc = Document()
    for sec in ["header"] + order:
        for kind, t in blocks[sec]:
            p = doc.add_paragraph()
            run = p.add_run(t.replace("- ", "• ", 1) if t.startswith("- ") else t)
            run.bold = kind in {"name", "h"}
    doc.save(path)


def write_pdf(path, blocks, order, two_col: bool):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.utils import simpleSplit
    from reportlab.pdfgen import canvas

    W, H = A4
    c = canvas.Canvas(str(path), pagesize=A4)

    def draw(x, y, items, width, size=9):
        for kind, t in items:
            font = "Helvetica-Bold" if kind in {"name", "h"} else "Helvetica"
            sz = 15 if kind == "name" else size
            if kind == "h":
                y -= 6
            text = ("• " + t[2:]) if t.startswith("- ") else t
            for piece in simpleSplit(text, font, sz, width):
                if y < 40:
                    c.showPage()
                    y = H - 50
                c.setFont(font, sz)
                c.drawString(x, y, piece)
                y -= sz + 3
        return y

    if not two_col:
        y = H - 50
        for sec in ["header"] + order:
            y = draw(40, y, blocks[sec], W - 80)
            y -= 4
    else:
        left_secs = ["header"] + [s for s in order if s in {"summary", "skills", "education"}]
        right_secs = [s for s in order if s in {"experience", "projects"}]
        y = H - 50
        for sec in left_secs:
            y = draw(40, y, blocks[sec], 225)
        y = H - 50
        for sec in right_secs:
            y = draw(305, y, blocks[sec], 250, size=8.5)
    c.save()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=40)
    ap.add_argument("--out", default="data/synthetic")
    ap.add_argument("--seed", type=int, default=42)
    args = ap.parse_args()
    rng = random.Random(args.seed)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    for i in range(args.n):
        layout = LAYOUTS[i % len(LAYOUTS)]
        gold, blocks, order = make_record(rng, layout)
        stem = f"syn_{i + 1:03d}_{layout}"
        if layout == "txt":
            (out / f"{stem}.txt").write_text(as_text(blocks, order), encoding="utf-8")
        elif layout == "docx":
            write_docx(out / f"{stem}.docx", blocks, order)
        else:
            write_pdf(out / f"{stem}.pdf", blocks, order, two_col=(layout == "pdf_two_col"))
        (out / f"{stem}.gold.json").write_text(json.dumps(gold, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"wrote {args.n} resumes + gold files to {out}")


if __name__ == "__main__":
    main()
